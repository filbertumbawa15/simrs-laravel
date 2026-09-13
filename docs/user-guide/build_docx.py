"""
Build DOCX user guide untuk SIHRS.

Usage:
    python build_docx.py <role>     # single role
    python build_docx.py all        # semua role
"""

import os
import sys
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE_DIR = Path(__file__).parent
SCREENSHOTS_DIR = BASE_DIR / 'screenshots'
OUT_DIR = BASE_DIR / 'docx-output'
OUT_DIR.mkdir(exist_ok=True)


# ============================================================
# Styling helpers
# ============================================================

def set_cell_bg(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), color_hex)
    tc_pr.append(shd)


def add_callout(doc, kind, text):
    colors = {
        'tip': ('E0F2FE', '075985', '💡 TIP'),
        'warning': ('FEF3C7', '92400E', '⚠️ PERHATIAN'),
        'critical': ('FEE2E2', '991B1B', '🚨 KRITIS'),
        'info': ('F3F4F6', '374151', 'ℹ️ INFO'),
    }
    bg, fg, label = colors.get(kind, colors['info'])
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_bg(cell, bg)
    p1 = cell.paragraphs[0]
    r1 = p1.add_run(label); r1.bold = True; r1.font.color.rgb = RGBColor.from_string(fg); r1.font.size = Pt(9)
    p2 = cell.add_paragraph()
    r2 = p2.add_run(text); r2.font.color.rgb = RGBColor.from_string(fg); r2.font.size = Pt(10)
    doc.add_paragraph()


def add_screenshot(doc, filename, caption, width_cm=15):
    path = SCREENSHOTS_DIR / filename
    if not path.exists():
        p = doc.add_paragraph()
        r = p.add_run(f'[Screenshot tidak tersedia: {filename} — capture manual saat rollout]')
        r.italic = True; r.font.color.rgb = RGBColor(180, 0, 0)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(); r.add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cap.add_run(f'Gambar: {caption}'); cr.italic = True; cr.font.size = Pt(9); cr.font.color.rgb = RGBColor(107, 114, 128)


def add_step(doc, number, title, body=None):
    p = doc.add_paragraph()
    r1 = p.add_run(f'Langkah {number} — '); r1.bold = True; r1.font.size = Pt(11); r1.font.color.rgb = RGBColor.from_string('0D9488')
    r2 = p.add_run(title); r2.bold = True; r2.font.size = Pt(11)
    if body: doc.add_paragraph(body)


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        r1 = p.add_run(bold_prefix); r1.bold = True
        p.add_run(text)
    else:
        p.add_run(text)


def add_cover(doc, title, subtitle, role_desc):
    t = doc.add_heading(title, level=0); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle); r.font.size = Pt(12); r.font.color.rgb = RGBColor(107, 114, 128)
    doc.add_paragraph(); doc.add_paragraph()
    add_callout(doc, 'info', role_desc)


def add_toc(doc, items):
    doc.add_heading('Daftar Isi', level=1)
    for t in items:
        p = doc.add_paragraph(t); p.paragraph_format.space_after = Pt(4)


def add_troubleshooting(doc, problems):
    doc.add_heading('Troubleshooting Umum', level=1)
    for problem, solution in problems:
        p = doc.add_paragraph()
        r = p.add_run(f'❌ {problem}'); r.bold = True
        doc.add_paragraph(f'→ {solution}')
        doc.add_paragraph()


def add_contact(doc, contacts):
    doc.add_heading('Kontak Bantuan', level=1)
    doc.add_paragraph('Isi kontak sesuai struktur RS Anda:')
    table = doc.add_table(rows=len(contacts) + 1, cols=2)
    table.style = 'Light Grid'
    table.rows[0].cells[0].text = 'Situasi'
    table.rows[0].cells[1].text = 'Kontak'
    for cell in table.rows[0].cells:
        for p in cell.paragraphs:
            for r in p.runs: r.bold = True
    for i, (situasi, kontak) in enumerate(contacts, start=1):
        table.rows[i].cells[0].text = situasi
        table.rows[i].cells[1].text = kontak
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('— Selamat bertugas! —'); r.italic = True; r.font.color.rgb = RGBColor(107, 114, 128)


def set_document_style(doc):
    styles = doc.styles
    normal = styles['Normal']; normal.font.name = 'Calibri'; normal.font.size = Pt(11)
    for lvl in [1, 2, 3]:
        h = styles[f'Heading {lvl}']; h.font.name = 'Calibri'; h.font.color.rgb = RGBColor.from_string('0D9488')


# ============================================================
# BUILDERS PER ROLE
# ============================================================

def build_getting_started(doc):
    add_cover(doc, 'Getting Started', 'Panduan Onboarding Universal untuk Semua Staff SIHRS',
              'Baca dokumen ini pertama kali sebelum panduan role spesifik Anda. Berisi cara login, mengenali UI, cetak dokumen, dan logout aman.')
    doc.add_page_break()

    add_toc(doc, ['1. Akses Aplikasi', '2. Login Pertama Kali', '3. Kenali Layout Aplikasi',
                  '4. Cara Membaca Status & Warna', '5. Tombol Umum', '6. Realtime Refresh',
                  '7. Cara Cetak Dokumen', '8. Logout Aman', '9. Yang Harus Dilakukan Kalau Ada Masalah'])
    doc.add_page_break()

    doc.add_heading('1. Akses Aplikasi', level=1)
    doc.add_paragraph('Buka browser (Chrome atau Edge disarankan) dan ketik alamat aplikasi yang diberikan oleh IT.')
    add_screenshot(doc, '00-halaman-login.png', 'Halaman login SIHRS.')
    add_callout(doc, 'critical', 'PASTIKAN URL diawali dengan https:// (ada gembok di address bar). Kalau tidak ada gembok, jangan login — hubungi IT.')

    doc.add_heading('2. Login Pertama Kali', level=1)
    add_step(doc, 1, 'Ketik Username', 'Username dari IT (contoh: dr.iqbal, apt.fitri, kasir.lina). Case-sensitive.')
    add_step(doc, 2, 'Ketik Password', 'Password sementara dari IT — WAJIB diganti setelah login pertama (hubungi IT untuk reset).')
    add_step(doc, 3, 'Klik "Masuk"')
    add_screenshot(doc, '00-form-login.png', 'Form login diisi username.')
    add_callout(doc, 'warning', 'Salah password 5x dalam 1 menit → akun ke-lock 1 menit. Fitur keamanan anti-hacker. Tunggu, coba lagi.')

    doc.add_page_break()
    doc.add_heading('3. Kenali Layout Aplikasi', level=1)
    doc.add_paragraph('Setelah login, Anda lihat 3 area:')
    add_bullet(doc, 'Sidebar kiri (hijau) — menu navigasi. Menu Anda BERBEDA dengan rekan role lain — bukan bug.', bold_prefix='A. ')
    add_bullet(doc, 'Top bar atas — jam realtime + nama Anda + dropdown logout.', bold_prefix='B. ')
    add_bullet(doc, 'Area konten utama — isi menu yang dipilih.', bold_prefix='C. ')
    add_screenshot(doc, '00-dashboard-layout.png', 'Dashboard SIHRS.')

    doc.add_heading('4. Cara Membaca Status & Warna', level=1)
    doc.add_paragraph('Warna badge/label punya arti konsisten di seluruh aplikasi:')
    add_bullet(doc, 'Hijau — Normal / selesai / OK (Lunas, Hasil Normal, Kamar Tersedia)', bold_prefix='🟢 ')
    add_bullet(doc, 'Biru — Info / dalam proses (Sedang Diperiksa)', bold_prefix='🔵 ')
    add_bullet(doc, 'Kuning / Oranye — Perhatian / menunggu (Menunggu Pembayaran, Stok Menipis)', bold_prefix='🟡 ')
    add_bullet(doc, 'Merah — Kritis / gagal (Hasil Kritis Lab, Kamar Penuh, Tagihan Belum Lunas)', bold_prefix='🔴 ')
    add_bullet(doc, 'Abu — Nonaktif / batal', bold_prefix='⚪ ')
    add_screenshot(doc, '00-status-badge.png', 'Contoh badge status di dashboard.')

    doc.add_heading('5. Tombol Umum', level=1)
    add_bullet(doc, 'Simpan/Store — simpan data. Setelah simpan, tidak semua field bisa diedit.')
    add_bullet(doc, 'Update/Perbarui — simpan perubahan, tercatat di audit log.')
    add_bullet(doc, 'Batal/Cancel — kembali tanpa simpan. Aman.')
    add_bullet(doc, 'Hapus/Delete — soft delete (data tidak hilang, hanya tidak tampak).')
    add_bullet(doc, 'Cetak / 🖨️ — buka PDF di tab baru (allow popup).')
    add_bullet(doc, 'Cari / Search — filter list, auto-update.')
    add_bullet(doc, 'Export Excel 📊 — download data ke .xlsx (role tertentu saja).')

    doc.add_page_break()
    doc.add_heading('6. Realtime Refresh (Board & Antrian)', level=1)
    doc.add_paragraph('Halaman berikut auto-refresh setiap 15–30 detik: Antrian RJ, IGD Board, Bed Management.')
    add_screenshot(doc, '00-auto-refresh-badge.png', 'Badge auto-refresh — pojok kanan atas, dengan tombol ⏸ pause.', width_cm=14)
    add_callout(doc, 'tip', 'Kalau sedang isi form lama, klik ⏸ dulu supaya halaman tidak refresh dan input hilang.')

    doc.add_heading('7. Cara Cetak Dokumen', level=1)
    add_step(doc, 1, 'Klik tombol 🖨️ atau Cetak di halaman terkait')
    add_step(doc, 2, 'Tab baru terbuka menampilkan PDF')
    add_step(doc, 3, 'Tekan Ctrl+P (atau Cmd+P) untuk print')
    add_step(doc, 4, 'Pilih printer tujuan (thermal 58mm untuk tiket, printer normal untuk kuitansi/resume)')
    add_screenshot(doc, '00-cetak-dialog.png', 'Halaman siap cetak.')
    add_callout(doc, 'warning', 'Setiap PDF ada QR Code verifikasi di footer. Jangan crop bagian ini saat print — QR bisa di-scan pasien untuk konfirmasi keaslian.')

    doc.add_heading('8. Logout Aman', level=1)
    add_callout(doc, 'critical', 'Jangan pernah tinggalkan komputer login. Semua aktivitas dicatat atas nama Anda di audit log — kalau orang lain pakai akun Anda, Anda yang bertanggung jawab.')
    add_step(doc, 1, 'Klik nama Anda di top-bar kanan atas')
    add_step(doc, 2, 'Klik menu "Logout"')
    add_screenshot(doc, '00-logout-menu.png', 'Dropdown user menu dengan opsi logout.')
    doc.add_paragraph('Session otomatis expired setelah 2 jam idle.')

    doc.add_page_break()
    doc.add_heading('9. Yang Harus Dilakukan Kalau Ada Masalah', level=1)
    doc.add_heading('Aplikasi lambat / hang', level=2)
    doc.add_paragraph('1. Tekan F5 refresh sekali. 2. Kalau tetap lambat >30 detik, buka /up/health di tab baru. 3. Kalau muncul JSON dengan "status": "ok" = sistem OK (masalah di komputer/network Anda). 4. Kalau tidak muncul = hubungi IT segera.')
    doc.add_heading('Ada pesan error merah', level=2)
    doc.add_paragraph('1. Baca pesannya — biasanya menjelaskan (mis. "Stok Paracetamol tidak cukup"). 2. Kalau teknis (kode/angka), screenshot dan kirim IT. 3. Jangan panik, jangan close browser — data terakhir mungkin masih tersimpan.')
    doc.add_heading('Data hilang / tidak muncul', level=2)
    doc.add_paragraph('1. Jangan re-entry dulu. Klik Reset filter dulu (data mungkin ter-hide oleh filter). 2. Kalau tetap tidak ada, hubungi IT — mereka bisa cek audit log.')

    add_contact(doc, [
        ('Aplikasi error / tidak bisa akses', 'IT On-call: __________'),
        ('Pertanyaan alur / SOP', 'Supervisor unit Anda: __________'),
        ('Data salah / butuh koreksi', 'IT Support: __________'),
    ])


# ============================================================
def build_registrasi(doc):
    add_cover(doc, 'Panduan Petugas Registrasi', 'SIHRS — Sistem Informasi Rumah Sakit',
              'Panduan ini untuk role REGISTRASI (contoh username: reg.mira). Baca dokumen "Getting Started" dulu sebelum melanjutkan.')
    doc.add_page_break()

    add_toc(doc, ['1. Tanggung Jawab Anda', '2. Skenario A: Pasien Baru Datang',
                  '3. Skenario B: Pasien Lama Datang Lagi', '4. Membuat Kunjungan',
                  '5. Mencetak Tiket Antrian', '6. Troubleshooting', '7. Kontak'])
    doc.add_page_break()

    doc.add_heading('1. Tanggung Jawab Anda', level=1)
    doc.add_paragraph('Tugas utama Anda:')
    add_bullet(doc, 'Mendaftarkan pasien baru yang belum pernah berobat.')
    add_bullet(doc, 'Mencari pasien lama yang datang lagi.')
    add_bullet(doc, 'Membuat kunjungan (RJ / IGD / RI).')
    add_bullet(doc, 'Mencetak tiket antrian.')
    doc.add_paragraph()
    p = doc.add_paragraph(); r = p.add_run('Yang TIDAK boleh Anda lakukan:'); r.bold = True
    add_bullet(doc, 'Menghapus data pasien (duplikat → laporkan IT).')
    add_bullet(doc, 'Melihat rekam medis / hasil lab / resep.')
    add_bullet(doc, 'Menerima pembayaran (tugas kasir).')

    doc.add_page_break()
    doc.add_heading('2. Skenario A: Pasien Baru Datang', level=1)
    add_callout(doc, 'critical', 'WAJIB cek dulu apakah pasien sudah pernah berobat sebelum daftar baru. Kalau tidak, akan muncul data duplikat.')

    add_step(doc, 1, 'Buka Menu Pasien', 'Klik "Pasien" di sidebar kiri.')
    add_screenshot(doc, '01-pasien-list-search.png', 'Daftar pasien dengan kolom pencarian.')

    add_step(doc, 2, 'Cari Pasien Berdasarkan NIK atau Nama', 'Ketik NIK (16 digit) atau nama lengkap di kolom pencarian.')
    add_callout(doc, 'tip', 'Cari pakai NIK lebih akurat daripada nama (nama bisa typo). Kalau anak tanpa NIK, kombinasi nama ibu + tgl lahir.')
    doc.add_paragraph('Kalau HASIL KOSONG → lanjut Langkah 3. Kalau HASIL DITEMUKAN → lompat ke Skenario B.')

    add_step(doc, 3, 'Klik Tombol "+ Pasien Baru"', 'Tombol di pojok kanan atas.')
    add_screenshot(doc, '01-form-pasien-baru.png', 'Form pendaftaran pasien baru.')

    add_step(doc, 4, 'Isi Data Identitas Pasien')
    doc.add_paragraph('Field wajib (bertanda *): Nama Lengkap, Jenis Kelamin, Tanggal Lahir, Alamat.')
    doc.add_paragraph('Field opsional tapi disarankan: NIK, No. Telp, Kontak Darurat.')

    add_step(doc, 5, 'Klik "Simpan & Buat Rekam Medis"', 'Sistem generate No. RM otomatis (contoh: 100051).')
    add_callout(doc, 'warning', 'Kalau NIK sudah dipakai pasien lain, sistem tolak. Cari dulu pasien yang punya NIK itu. Kalau orang beda, screenshot dan kirim IT.')

    doc.add_page_break()
    doc.add_heading('3. Skenario B: Pasien Lama Datang Lagi', level=1)
    doc.add_paragraph('Kalau saat pencarian pasien sudah ada, jangan daftarkan ulang. Langsung buat kunjungan.')
    add_step(doc, 1, 'Klik Nama Pasien di Hasil Pencarian')
    add_screenshot(doc, '01-pasien-detail-with-riwayat.png', 'Detail pasien dengan riwayat kunjungan.')
    add_step(doc, 2, 'Klik "+ Kunjungan Baru"', 'Tombol di pojok kanan atas halaman detail.')

    doc.add_page_break()
    doc.add_heading('4. Membuat Kunjungan', level=1)
    add_screenshot(doc, '01-form-kunjungan.png', 'Form kunjungan baru.')
    doc.add_paragraph('Isi form:')
    p = doc.add_paragraph(); r = p.add_run('Tipe Kunjungan:'); r.bold = True
    add_bullet(doc, 'RJ (Rawat Jalan) — paling umum, konsultasi poli.')
    add_bullet(doc, 'IGD — pasien gawat darurat.')
    add_bullet(doc, 'RI (Rawat Inap) — biasanya dari IGD/poli, bukan langsung dari registrasi.')
    p = doc.add_paragraph(); r = p.add_run('Poli & Dokter:'); r.bold = True
    doc.add_paragraph('Hanya muncul kalau tipe RJ. Pilih poli tujuan dan dokter yang praktik hari ini.')
    p = doc.add_paragraph(); r = p.add_run('Penjamin:'); r.bold = True
    add_bullet(doc, 'UMUM — bayar sendiri.')
    add_bullet(doc, 'BPJS — isi No. SEP dari kartu.')
    add_bullet(doc, 'ASURANSI — pilih asuransi kerjasama.')

    add_callout(doc, 'warning', 'Kalau pasien masih punya kunjungan aktif (belum SELESAI), sistem tolak. Selesaikan kunjungan lama dulu.')
    add_step(doc, 1, 'Klik Simpan', 'No. Kunjungan otomatis generated (contoh: RJ/2026/09/00042).')
    add_screenshot(doc, '01-kunjungan-berhasil.png', 'Detail kunjungan setelah dibuat.')

    doc.add_page_break()
    doc.add_heading('5. Mencetak Tiket Antrian', level=1)
    add_step(doc, 1, 'Buka Menu Antrian RJ', 'Cari nama pasien di list.')
    add_step(doc, 2, 'Klik Tombol 🖨️', 'Tab baru terbuka menampilkan PDF tiket antrian.')
    add_step(doc, 3, 'Print ke Printer Thermal', 'Ctrl+P → pilih printer thermal 58mm → Print.')
    add_callout(doc, 'tip', 'Kalau printer thermal 80mm, cetak di kertas A5 pakai printer biasa juga terbaca.')

    add_troubleshooting(doc, [
        ('Form daftar pasien error "NIK sudah terdaftar"',
         'Pasien sudah pernah datang. Cari di daftar pasien pakai NIK tsb. Kalau orang beda dengan NIK sama, screenshot dan kirim IT.'),
        ('Tombol Simpan tidak reaksi',
         'Ada field wajib (*) yang belum diisi. Scroll ke atas, cek label merah.'),
        ('Muncul pesan "Terlalu banyak percobaan"',
         'Rate limit terpicu (30 request/menit). Tunggu 1 menit. Cukup klik sekali, jangan double-click.'),
        ('Pasien terdaftar dua kali (duplikat)',
         'Jangan hapus. Screenshot kedua record, kirim IT untuk merge. Sementara pakai No. RM lebih kecil.'),
        ('Kunjungan salah dibuat',
         'Selama status TERDAFTAR, buka detail → Batal. Kalau sudah DALAM_PEMERIKSAAN, hubungi IT.'),
    ])

    add_contact(doc, [
        ('Aplikasi error', 'IT On-call: __________'),
        ('Pertanyaan alur/SOP', 'Supervisor Pendaftaran: __________'),
        ('Data koreksi', 'IT Support: __________'),
    ])


# ============================================================
def build_dokter(doc):
    add_cover(doc, 'Panduan Dokter', 'DOKTER Umum & DOKTER Spesialis',
              'Panduan ini untuk role DOKTER (dr.iqbal) & DOKTER_SPESIALIS (dr.andika/dr.sari). Spesialis punya akses tambahan untuk admisi & discharge Rawat Inap sebagai DPJP.')
    doc.add_page_break()

    add_toc(doc, ['1. Perbedaan Role DOKTER vs DOKTER_SPESIALIS',
                  '2. Alur Kerja Rawat Jalan (RJ)',
                  '3. Isi SOAP + Tanda Vital',
                  '4. Isi Diagnosa (ICD-10 Wajib)',
                  '5. Order Lab / Radiologi',
                  '6. Buat Resep',
                  '7. Selesaikan Pemeriksaan',
                  '8. Rawat Inap: Admisi (khusus DPJP)',
                  '9. Rawat Inap: Discharge (khusus DPJP)',
                  '10. Melihat Hasil Lab / Radiologi',
                  '11. Troubleshooting', '12. Kontak'])
    doc.add_page_break()

    doc.add_heading('1. Perbedaan Role', level=1)
    p = doc.add_paragraph()
    p.add_run('DOKTER — ').bold = True
    p.add_run('Periksa RJ/IGD, order lab/rad, buat resep. Tidak bisa admisi/discharge RI.')
    p = doc.add_paragraph()
    p.add_run('DOKTER_SPESIALIS — ').bold = True
    p.add_run('Semua di atas + bisa jadi DPJP di RI (admisi, CPPT, pindah kamar, discharge).')

    doc.add_page_break()
    doc.add_heading('2. Alur Kerja Rawat Jalan', level=1)
    add_step(doc, 1, 'Buka Antrian Poli Anda', 'Login → sidebar → menu "Antrian RJ". Filter berdasarkan poli Anda.')
    add_screenshot(doc, '02-antrian-rj-list.png', 'Halaman antrian RJ semua poli.')
    add_callout(doc, 'tip', 'Bookmark URL antrian poli Anda untuk akses cepat.')

    add_step(doc, 2, 'Panggil Pasien', 'Cari pasien di baris antrian. Klik tombol 📢 Panggil. Sistem catat waktu panggilan.')
    add_screenshot(doc, '02-antrian-tombol-panggil.png', 'Baris antrian dengan tombol Panggil dan Mulai Periksa.')

    add_step(doc, 3, 'Mulai Pemeriksaan', 'Klik "Mulai Periksa" saat pasien masuk ruangan. Halaman SOAP terbuka.')
    add_screenshot(doc, '02-form-soap-empty.png', 'Form SOAP kosong dengan sidebar informasi pasien.')

    doc.add_page_break()
    doc.add_heading('3. Isi SOAP + Tanda Vital', level=1)
    p = doc.add_paragraph(); r = p.add_run('S (Subjective) — '); r.bold = True
    p.add_run('Keluhan pasien, anamnesa (contoh: "Nyeri kepala sejak 3 hari").')
    p = doc.add_paragraph(); r = p.add_run('O (Objective) — '); r.bold = True
    p.add_run('Pemeriksaan fisik + tanda vital (TD, N, R, S, SpO2, BB, TB).')
    p = doc.add_paragraph(); r = p.add_run('A (Assessment) — '); r.bold = True
    p.add_run('Penilaian klinis selain diagnosa.')
    p = doc.add_paragraph(); r = p.add_run('P (Plan) — '); r.bold = True
    p.add_run('Rencana terapi: obat, follow-up, edukasi.')
    add_screenshot(doc, '02-form-soap-diisi.png', 'Form SOAP dengan sample data terisi.')

    doc.add_page_break()
    doc.add_heading('4. Isi Diagnosa (ICD-10 Wajib)', level=1)
    add_step(doc, 1, 'Scroll ke Section Diagnosa', 'Klik "+ Tambah Diagnosa".')
    add_step(doc, 2, 'Cari ICD-10', 'Ketik nama penyakit atau kode → autocomplete muncul.')
    add_screenshot(doc, '02-diagnosa-search-icd.png', 'Autocomplete ICD-10 saat mengetik.')
    add_step(doc, 3, 'Pilih Tipe', 'PRIMER (utama) / SEKUNDER / KOMPLIKASI.')
    add_step(doc, 4, 'Klik Simpan')
    add_callout(doc, 'critical', 'WAJIB minimum 1 diagnosa PRIMER. Sistem tolak "Selesai" kalau tidak ada.')

    doc.add_heading('5. Order Lab / Radiologi', level=1)
    doc.add_paragraph('Dari halaman detail kunjungan atau menu Lab/Radiologi, klik "+ Buat Order".')
    add_screenshot(doc, '02-order-lab-form.png', 'Form order lab dengan multi-select parameter.')
    doc.add_paragraph('Isi: parameter yang diminta, prioritas (RUTIN/CITO), catatan klinis, diagnosa kerja.')
    add_callout(doc, 'critical', 'Untuk radiologi X-Ray/CT pada wanita usia subur, WAJIB centang "Hamil?" dulu. Kalau lupa, radiografer bisa lakukan X-Ray ke pasien hamil.')

    doc.add_page_break()
    doc.add_heading('6. Buat Resep', level=1)
    add_screenshot(doc, '02-form-resep.png', 'Form resep dengan multi obat.')
    doc.add_paragraph('Per obat: ketik nama (autocomplete + info stok), jumlah, signa (3x1, 2x1, k/p), aturan pakai, catatan.')
    add_callout(doc, 'tip', 'Kalau obat racikan (puyer), centang "is_racikan".')

    doc.add_heading('7. Selesaikan Pemeriksaan', level=1)
    doc.add_paragraph('Setelah semua urusan beres (SOAP, diagnosa, order, resep), klik "Selesai".')
    doc.add_paragraph('Sistem cek: ada diagnosa PRIMER? OK → status pasien auto-update:')
    add_bullet(doc, 'Ada order lab → MENUNGGU_HASIL_LAB')
    add_bullet(doc, 'Ada resep → MENUNGGU_OBAT')
    add_bullet(doc, 'Tidak ada → MENUNGGU_PEMBAYARAN')
    add_screenshot(doc, '02-selesai-berhasil.png', 'Halaman antrian setelah pasien selesai (pasien hilang dari list).')

    doc.add_page_break()
    doc.add_heading('8. Rawat Inap: Admisi (khusus DPJP)', level=1)
    add_step(doc, 1, 'Buka Menu Rawat Inap → Admisi Baru')
    add_screenshot(doc, '02-form-admisi-ri.png', 'Form admisi rawat inap.')
    add_step(doc, 2, 'Pilih Kamar', 'Dropdown menampilkan kamar TERSEDIA + kelas + tarif.')
    add_step(doc, 3, 'Isi Alasan Masuk', 'Indikasi rawat inap yang jelas (contoh: "Dehidrasi berat, perlu observasi + rehidrasi IV").')
    add_step(doc, 4, 'Klik Simpan', 'Kamar auto-set jadi TERISI. Sistem generate No. Kunjungan RI.')

    doc.add_heading('9. Rawat Inap: Discharge (khusus DPJP)', level=1)
    add_step(doc, 1, 'Halaman Detail RI → Klik "Pulangkan Pasien"')
    add_step(doc, 2, 'Pilih Cara Pulang',
             'SEMBUH / MEMBAIK / BELUM_SEMBUH / APS (Atas Permintaan Sendiri) / RUJUK / MENINGGAL.')
    add_step(doc, 3, 'Isi Resume Medis', 'WAJIB minimum 50 karakter — ringkasan perawatan.')
    add_step(doc, 4, 'Isi Instruksi Pulang', 'Obat, kontrol, restriksi aktivitas.')
    add_step(doc, 5, 'Klik "Simpan → Finalize Resume"')
    add_screenshot(doc, '02-form-discharge.png', 'Form discharge rawat inap.')
    add_callout(doc, 'critical', 'Resume medis TIDAK BISA di-edit setelah finalize. Baca ulang sebelum klik.')

    doc.add_page_break()
    doc.add_heading('10. Melihat Hasil Lab / Radiologi', level=1)
    doc.add_paragraph('Ketika Anda order lab/rad dan hasil sudah divalidasi:')
    add_bullet(doc, 'Notifikasi email kalau ada hasil kritis (flag LL/HH atau temuan kritis radiologi).')
    add_bullet(doc, 'Lihat manual di halaman detail kunjungan pasien.')
    add_bullet(doc, 'Cetak PDF hasil lab: menu PDF → Hasil Lab.')
    add_screenshot(doc, '02-detail-kunjungan-hasil-lab.png', 'Detail kunjungan dengan hasil lab tervalidasi.')
    add_callout(doc, 'critical', 'Email "🚨 [KRITIS] Hasil Lab" harus segera dilihat dan ditindaklanjuti. Email dikirim sesaat setelah dokter PK validasi.')

    add_troubleshooting(doc, [
        ('Autocomplete ICD-10 tidak muncul',
         'Ketik terlalu cepat (rate limit 120/menit). Jeda sebentar.'),
        ('"Selesai" gagal - diagnosa primer',
         'Tambahkan minimum 1 diagnosa dengan tipe PRIMER.'),
        ('"Simpan Resep" error stok tidak cukup',
         'Tetap bisa simpan resep — cek jumlah atau ganti obat. Kalau tetap ingin resep obat kosong (pasien beli luar), turunkan jumlah dan tulis catatan.'),
        ('Antrian tidak refresh sendiri',
         'Cek badge auto-refresh kanan atas. Kalau ⏸ = paused, klik ▶ untuk resume.'),
    ])

    add_contact(doc, [
        ('Konsultasi medis', 'Ketua Komdik: __________'),
        ('Bug aplikasi', 'IT Support: __________'),
    ])


# ============================================================
def build_perawat(doc):
    add_cover(doc, 'Panduan Perawat', 'PERAWAT Umum & PERAWAT Rawat Inap',
              'Panduan untuk role PERAWAT (pwt.rina) dan PERAWAT_RI (pwt.yusuf). PERAWAT_RI punya akses tambahan CPPT, request pindah kamar, dan update status kamar.')
    doc.add_page_break()

    add_toc(doc, ['1. Perbedaan Role', '2. Alur Bantuan Dokter di Poli RJ',
                  '3. Alur Rawat Inap (PERAWAT_RI)', '4. CPPT Keperawatan',
                  '5. Update Status Kamar', '6. Lihat Riwayat Medis Pasien',
                  '7. Triase IGD (perawat jaga)', '8. Troubleshooting', '9. Kontak'])
    doc.add_page_break()

    doc.add_heading('1. Perbedaan Role', level=1)
    p = doc.add_paragraph(); p.add_run('PERAWAT — ').bold = True
    p.add_run('Assist dokter periksa, cek/input tanda vital, view data pasien, triase IGD.')
    p = doc.add_paragraph(); p.add_run('PERAWAT_RI — ').bold = True
    p.add_run('Semua di atas + isi CPPT keperawatan, request pindah kamar, update status kamar.')

    doc.add_heading('2. Alur Bantuan Dokter di Poli RJ', level=1)
    add_step(doc, 1, 'Panggil Pasien', 'Koordinasi dengan dokter poli. Tombol Panggil di antrian.')
    add_step(doc, 2, 'Ukur Tanda Vital',
             'TD, Nadi, Respirasi, Suhu, SpO2, BB, TB. Input di form SOAP saat dokter buka pemeriksaan.')
    add_screenshot(doc, '03-tanda-vital-form.png', 'Bagian tanda vital di form SOAP.')
    add_callout(doc, 'tip', 'Kalau pasien sering kontrol, tanda vital sebelumnya tampak di sidebar untuk trend.')

    doc.add_page_break()
    doc.add_heading('3. Alur Rawat Inap (PERAWAT_RI)', level=1)
    add_step(doc, 1, 'Buka Daftar Pasien Rawat Inap', 'Menu "Rawat Inap".')
    add_screenshot(doc, '03-ri-list.png', 'Daftar pasien rawat inap.')

    add_step(doc, 2, 'Lihat Bed Management', 'Menu "Bed Management" — visualisasi status semua bed.')
    add_screenshot(doc, '03-bed-board.png', 'Bed management board.')
    doc.add_paragraph('Warna:')
    add_bullet(doc, 'Hijau = TERSEDIA', bold_prefix='🟢 ')
    add_bullet(doc, 'Merah = TERISI', bold_prefix='🔴 ')
    add_bullet(doc, 'Kuning = KOTOR (perlu cleaning)', bold_prefix='🟡 ')
    add_bullet(doc, 'Abu = MAINTENANCE (rusak)', bold_prefix='⚪ ')

    doc.add_page_break()
    doc.add_heading('4. CPPT Keperawatan', level=1)
    doc.add_paragraph('Setiap shift (pagi/siang/malam) WAJIB isi CPPT.')
    add_step(doc, 1, 'Menu Rawat Inap → Cari Pasien → Detail RI')
    add_step(doc, 2, 'Section CPPT → Klik "+ Tambah CPPT"')
    add_step(doc, 3, 'Pilih Profesi PERAWAT + Isi SOAP')
    add_screenshot(doc, '03-cppt-form.png', 'Form CPPT keperawatan.')
    doc.add_paragraph('Isi:')
    add_bullet(doc, 'Subjective: keluhan pasien (mis. "Nyeri area operasi skala 4/10")')
    add_bullet(doc, 'Objective: observasi (mis. "Luka tertutup rapi, tidak ada rembesan")')
    add_bullet(doc, 'Assessment: penilaian keperawatan')
    add_bullet(doc, 'Plan: rencana asuhan')
    add_bullet(doc, 'Instruksi: dari DPJP kalau ada')
    add_callout(doc, 'tip', 'CPPT semua profesi (dokter, perawat, apoteker, gizi) di 1 timeline urut waktu.')

    doc.add_heading('5. Update Status Kamar', level=1)
    doc.add_paragraph('Setelah pasien pulang, kamar auto-set KOTOR. Setelah cleaning selesai:')
    add_step(doc, 1, 'Menu Bed Management → Cari Kamar Kuning')
    add_step(doc, 2, 'Klik Kamar → Tombol "Ubah Status ke TERSEDIA"')
    add_screenshot(doc, '03-update-status-kamar.png', 'Ubah status kamar.')

    doc.add_heading('6. Lihat Riwayat Medis Pasien', level=1)
    add_screenshot(doc, '03-pasien-detail-riwayat.png', 'Detail pasien dengan riwayat.')
    add_callout(doc, 'critical', 'SELALU cek section "Alergi Obat" sebelum berikan obat apapun. Kalau ada alergi, jangan berikan — lapor DPJP.')

    doc.add_heading('7. Triase IGD (Perawat Jaga)', level=1)
    doc.add_paragraph('Kalau bertugas di IGD, Anda juga triase. Lihat panduan IGD terpisah.')

    add_troubleshooting(doc, [
        ('Bed Management tidak update setelah pasien pulang', 'Refresh (F5). Kalau tetap, hubungi IT.'),
        ('Tidak bisa klik "Tambah CPPT"', 'Cek role Anda — PERAWAT saja tidak bisa CPPT RI. Perlu PERAWAT_RI.'),
        ('Tanda vital anak, tidak ada lingkar kepala', 'v1 belum ada. Tulis di Objective PE.'),
    ])

    add_contact(doc, [
        ('Asuhan keperawatan', 'Kepala Ruangan: __________'),
        ('Bug aplikasi', 'IT Support: __________'),
    ])


# ============================================================
def build_lab(doc):
    add_cover(doc, 'Panduan Laboratorium', 'ANALIS_LAB & DOKTER_PK',
              'Panduan untuk role ANALIS_LAB (lab.budi) — sampling & input hasil, dan DOKTER_PK (dr.doni) — validasi hasil. Prinsip: separation of duties — orang yang input hasil tidak boleh sama dengan yang validasi.')
    doc.add_page_break()

    add_toc(doc, ['1. Perbedaan Role', '2. Alur Analis: Sampling & Input',
                  '3. Input Hasil + Auto-Flag', '4. Alur Dokter PK: Review & Validasi',
                  '5. Cetak PDF Hasil Lab', '6. Notifikasi Hasil Kritis',
                  '7. Troubleshooting', '8. Kontak'])
    doc.add_page_break()

    doc.add_heading('1. Perbedaan Role', level=1)
    p = doc.add_paragraph(); p.add_run('ANALIS_LAB — ').bold = True
    p.add_run('Sampling (mark sampel diambil), proses sampel, input hasil. TIDAK bisa validate.')
    p = doc.add_paragraph(); p.add_run('DOKTER_PK — ').bold = True
    p.add_run('Review dan validasi hasil. Hanya PK yang bisa finalize/release hasil.')

    doc.add_heading('2. Alur Analis: Sampling & Input', level=1)
    add_step(doc, 1, 'Buka Worklist', 'Menu "Laboratorium".')
    add_screenshot(doc, '04-lab-worklist.png', 'Worklist lab dengan filter status & prioritas.')
    add_callout(doc, 'tip', 'Sort by prioritas → CITO (merah) di atas. Kerjakan CITO dulu.')

    add_step(doc, 2, 'Sampling', 'Ketika pasien datang → cek detail order → ambil sampel → klik "Sampel Diambil".')
    add_screenshot(doc, '04-detail-order-lab.png', 'Detail order lab dengan list parameter.')

    add_step(doc, 3, 'Mulai Proses', 'Load sampel ke analyzer / tes manual. Klik "▶️ Mulai Proses".')

    doc.add_page_break()
    doc.add_heading('3. Input Hasil + Auto-Flag', level=1)
    add_step(doc, 1, 'Klik "📝 Input Hasil"')
    add_screenshot(doc, '04-form-input-hasil.png', 'Form input hasil lab.')
    doc.add_paragraph('Per parameter: nilai hasil (angka/teks), satuan auto-fill, catatan opsional.')
    doc.add_paragraph('Sistem auto-flag:')
    add_bullet(doc, 'N (Normal) — dalam range rujukan', bold_prefix='🟢 ')
    add_bullet(doc, 'L / H — di bawah / atas normal (bukan kritis)', bold_prefix='🟡 ')
    add_bullet(doc, 'LL / HH — KRITIS (nilai low/high kritis)', bold_prefix='🔴 ')
    add_screenshot(doc, '04-hasil-dengan-flag-kritis.png', 'Hasil dengan flag HH ter-highlight merah.')
    add_callout(doc, 'critical', 'Kalau flag LL/HH, WAJIB segera lapor telpon ke DPJP (jangan andalkan email saja).')
    add_callout(doc, 'warning', 'Jangan ubah hasil kalau tidak sesuai output analyzer. Kalau curiga error alat, ulang test.')

    add_step(doc, 2, 'Klik "Simpan Hasil"', 'Status order → VALIDASI (menunggu dokter PK).')

    doc.add_page_break()
    doc.add_heading('4. Alur Dokter PK: Review & Validasi', level=1)
    add_step(doc, 1, 'Filter Worklist Status = VALIDASI')
    add_screenshot(doc, '04-worklist-validasi.png', 'Worklist filter validasi.')

    add_step(doc, 2, 'Review Semua Parameter',
             'Cek apakah nilai masuk akal secara klinis. Sesuai diagnosa kerja? Ada flag kritis?')
    add_screenshot(doc, '04-review-hasil-lengkap.png', 'Halaman review hasil sebelum validate.')
    add_callout(doc, 'warning', 'Jangan validate tanpa baca. Nama Anda tercantum di PDF hasil — Anda bertanggung jawab.')

    add_step(doc, 3, 'Klik "✅ Validasi Hasil"', 'Status → SELESAI. Email peringatan ke DPJP terkirim otomatis kalau ada flag kritis.')
    add_screenshot(doc, '04-hasil-tervalidasi.png', 'Halaman hasil setelah divalidasi.')

    doc.add_heading('5. Cetak PDF Hasil Lab', level=1)
    doc.add_paragraph('Setelah divalidasi, klik "🖨️ Cetak Hasil PDF". PDF berisi kop RS, data pasien, semua parameter + flag, tanda tangan validator, QR verifikasi.')
    add_screenshot(doc, '04-pdf-hasil-lab.png', 'Preview PDF hasil lab.')

    doc.add_heading('6. Notifikasi Hasil Kritis (Otomatis)', level=1)
    doc.add_paragraph('Ketika PK validasi hasil dengan flag LL/HH, sistem otomatis:')
    add_bullet(doc, 'Kirim email peringatan ke email DPJP')
    add_bullet(doc, 'CC ke Komdik/Manajer Pelayanan (kalau di-set)')
    add_bullet(doc, 'Log ke audit trail (retensi 5 tahun)')
    add_callout(doc, 'warning', 'Email adalah backup, BUKAN pengganti komunikasi lisan. Untuk kritis, wajib juga telpon DPJP.')

    add_troubleshooting(doc, [
        ('Analyzer error, hasil korup', 'Jangan simpan hasil salah. Ulang test kalau sampel cukup.'),
        ('Pasien sampling tanpa order', 'Jangan sampling tanpa order. Minta dokter buat order dulu.'),
        ('Bisa cetak PDF hasil belum divalidasi?', 'Tidak. Sistem tolak sebelum status SELESAI.'),
        ('Dokter PK libur, siapa validate?', 'Wajib ada backup (RS lain kerjasama / internist mumpuni). Hubungi IT untuk akses temporer.'),
    ])

    add_contact(doc, [
        ('Konsultasi teknis lab', 'Ketua Instalasi Lab: __________'),
        ('Alat rusak', 'Vendor Analyzer: __________'),
        ('Bug aplikasi', 'IT Support: __________'),
    ])


# ============================================================
def build_radiologi(doc):
    add_cover(doc, 'Panduan Radiologi', 'RADIOGRAFER & DOKTER_RADIOLOG',
              'Panduan untuk role RADIOGRAFER (rad.tono) — eksekusi pemeriksaan & upload gambar, dan DOKTER_RADIOLOG (dr.dewi) — bacaan & validasi. Separation of duties berlaku.')
    doc.add_page_break()

    add_toc(doc, ['1. Perbedaan Role', '2. Alur Radiografer: Eksekusi',
                  '3. Upload Image', '4. Alur Radiolog: Bacaan',
                  '5. Validasi Order', '6. Troubleshooting', '7. Kontak'])
    doc.add_page_break()

    doc.add_heading('1. Perbedaan Role', level=1)
    p = doc.add_paragraph(); p.add_run('RADIOGRAFER — ').bold = True
    p.add_run('Eksekusi pemeriksaan, upload image. TIDAK bisa baca hasil.')
    p = doc.add_paragraph(); p.add_run('DOKTER_RADIOLOG — ').bold = True
    p.add_run('Baca gambar, input interpretasi (bacaan/kesan/saran), validasi order.')

    doc.add_heading('2. Alur Radiografer: Eksekusi', level=1)
    add_step(doc, 1, 'Buka Worklist', 'Menu "Radiologi".')
    add_screenshot(doc, '05-rad-worklist.png', 'Worklist radiologi.')
    add_callout(doc, 'critical',
        'SEBELUM eksekusi X-Ray/CT pada pasien wanita usia subur, WAJIB cek kolom "Hamil" di detail order. Kalau ada tanda ⚠️ Hamil, JANGAN lakukan X-Ray tanpa konfirmasi ulang ke DPJP.')

    add_step(doc, 2, 'Cek Persiapan Pasien', 'Buka detail order untuk cek: jenis pemeriksaan, keterangan klinis, puasa/hamil.')
    add_screenshot(doc, '05-detail-order-rad.png', 'Detail order radiologi.')

    add_step(doc, 3, 'Konfirmasi Identitas Pasien',
             'Nama + tgl lahir (jangan hanya nama). Untuk wanita: konfirmasi ulang tidak hamil / siklus haid terakhir.')

    doc.add_page_break()
    doc.add_heading('3. Upload Image', level=1)
    add_step(doc, 1, 'Klik "📸 Eksekusi Pemeriksaan"')
    add_screenshot(doc, '05-form-eksekusi.png', 'Form eksekusi dengan upload image.')

    add_step(doc, 2, 'Isi Kondisi Teknis',
             'Parameter foto (contoh: "Thorax PA, kV 65, mAs 8, jarak 150cm, pasien inspirasi maksimal").')

    add_step(doc, 3, 'Upload Image',
             'Format: JPG, PNG, PDF, DCM. Max 20 MB per file.')

    add_step(doc, 4, 'Klik "Upload & Simpan"', 'Status → MENUNGGU_BACAAN.')
    add_screenshot(doc, '05-eksekusi-selesai.png', 'Order dengan images ter-upload.')
    add_screenshot(doc, '05-viewer-image.png', 'Preview image radiologi.')

    doc.add_page_break()
    doc.add_heading('4. Alur Radiolog: Bacaan', level=1)
    add_step(doc, 1, 'Filter Worklist Status = MENUNGGU_BACAAN')
    add_step(doc, 2, 'Lihat Gambar', 'Klik image untuk full-size di tab baru. Untuk DICOM, buka pakai viewer khusus (RadiAnt, OsiriX).')
    add_step(doc, 3, 'Klik "📝 Input Bacaan"')
    add_screenshot(doc, '05-form-bacaan.png', 'Form input bacaan radiologi.')

    doc.add_paragraph('Isi:')
    add_bullet(doc, 'Bacaan: deskripsi radiologis lengkap (mis. "Cor: CTR 55%. Pulmo: infiltrat lobus inferior kanan berbatas tidak tegas.")')
    add_bullet(doc, 'Kesan: kesimpulan singkat (mis. "Pneumonia lobar dekstra.")')
    add_bullet(doc, 'Saran: follow-up disarankan')
    add_bullet(doc, 'Ada Temuan Kritis? — CENTANG jika ada (pneumothorax, perdarahan intracranial, fraktur mayor, dll)')
    add_callout(doc, 'critical', 'Centang "Ada Temuan Kritis" dengan hati-hati — memicu email peringatan otomatis ke DPJP. False positive = spam DPJP, false negative = pasien terlambat ditangani.')

    doc.add_heading('5. Validasi Order', level=1)
    add_step(doc, 1, 'Klik "✅ Validasi Order"')
    doc.add_paragraph('Sistem cek: semua pemeriksaan sudah ada bacaan? Ada minimum 1 image? Kalau OK → status SELESAI, bacaan release.')
    add_screenshot(doc, '05-validasi-berhasil.png', 'Order status SELESAI.')
    add_callout(doc, 'warning', 'Setelah validasi, bacaan TIDAK BISA diedit. Baca ulang sebelum klik.')

    add_troubleshooting(doc, [
        ('Foto blur, bisa ulang?', 'Ya. Upload versi baru. Hapus yang blur sebelum validasi.'),
        ('Pasien walk-in tanpa order', 'Radiografer TIDAK BOLEH eksekusi tanpa order. Minta dokter buat order dulu.'),
        ('Upload image gagal ("File terlalu besar")', 'Max 20 MB. Kompres pakai tinyjpg.com atau tools DICOM converter.'),
        ('"Validasi" error "Belum semua dibaca"', 'Cek order multi-pemeriksaan — semua harus punya bacaan.'),
    ])

    add_contact(doc, [
        ('Konsultasi teknis', 'Ketua Instalasi Radiologi: __________'),
        ('Alat rusak', 'Vendor Alat: __________'),
        ('Bug aplikasi', 'IT Support: __________'),
    ])


# ============================================================
def build_farmasi(doc):
    add_cover(doc, 'Panduan Farmasi (Apotek)', 'APOTEKER & TTK_APOTEK',
              'Panduan untuk role APOTEKER (apt.fitri) — verifikasi & dispensing, dan TTK_APOTEK (tta.dewi) — dispensing saja di bawah supervisi. Verifikasi wajib oleh apoteker berlisensi.')
    doc.add_page_break()

    add_toc(doc, ['1. Perbedaan Role', '2. Buka Worklist Resep',
                  '3. Verifikasi Resep (APOTEKER)', '4. Serahkan Obat (Dispense)',
                  '5. Cetak Resep', '6. Kelola Stok', '7. Troubleshooting', '8. Kontak'])
    doc.add_page_break()

    doc.add_heading('1. Perbedaan Role', level=1)
    p = doc.add_paragraph(); p.add_run('APOTEKER — ').bold = True
    p.add_run('Verifikasi resep (cek dosis, interaksi, alergi) + dispensing.')
    p = doc.add_paragraph(); p.add_run('TTK_APOTEK — ').bold = True
    p.add_run('Dispensing saja (di bawah supervisi apoteker). TIDAK bisa verifikasi.')

    doc.add_heading('2. Buka Worklist Resep', level=1)
    doc.add_paragraph('Menu "Farmasi" — worklist resep aktif.')
    add_screenshot(doc, '06-resep-worklist.png', 'Worklist resep dengan filter status.')
    doc.add_paragraph('Filter status: BARU → DIVERIFIKASI → DISERAHKAN → BATAL')

    doc.add_page_break()
    doc.add_heading('3. Verifikasi Resep (APOTEKER)', level=1)
    add_step(doc, 1, 'Klik Detail Resep Status BARU')
    add_screenshot(doc, '06-detail-resep-verify.png', 'Detail resep sebelum verifikasi.')

    p = doc.add_paragraph(); r = p.add_run('Yang harus dicek:'); r.bold = True
    add_bullet(doc, 'Alergi: cek section alergi obat di data pasien.')
    add_bullet(doc, 'Dosis: sesuai umur/BB pasien?')
    add_bullet(doc, 'Interaksi: ada interaksi obat berbahaya?')
    add_bullet(doc, 'Frekuensi: signa masuk akal?')
    add_bullet(doc, 'Duplikasi: obat sama dengan mekanisme aksi mirip?')
    add_bullet(doc, 'Kontraindikasi: sesuai kondisi (mis. NSAID untuk ulkus?)')

    add_step(doc, 2, 'Kalau Semua OK → Klik "✅ Verifikasi Resep"', 'Status → DIVERIFIKASI. Nama Anda tercatat sebagai verifikator.')
    p = doc.add_paragraph(); r = p.add_run('Kalau ada masalah:'); r.bold = True
    doc.add_paragraph('Jangan verifikasi. Telpon dokter, diskusi, minta koreksi. Dokter edit/buat resep baru.')
    add_callout(doc, 'critical', 'Verifikasi tanpa cek = malpraktik farmasi. Nama Anda di sistem sebagai verifikator — Anda bertanggung jawab.')

    doc.add_page_break()
    doc.add_heading('4. Serahkan Obat (Dispense)', level=1)
    add_step(doc, 1, 'Cari Resep Status DIVERIFIKASI')
    add_step(doc, 2, 'Ambil Obat dari Rak', 'Sistem otomatis pilih batch FEFO (First Expired First Out).')
    add_screenshot(doc, '06-detail-resep-dispense.png', 'Detail resep siap dispense.')

    add_step(doc, 3, 'Klik "📦 Serahkan Obat"', 'Sistem cek stok cukup untuk semua obat. Kalau ada 1 kurang → transaksi rollback.')
    add_callout(doc, 'warning', 'Sistem TIDAK bisa dispense parsial. Semua obat harus tersedia. Kalau 1 kurang, telpon dokter untuk resep alternatif atau catat pasien beli di apotek luar.')

    add_screenshot(doc, '06-serahkan-success.png', 'Resep status DISERAHKAN dengan batch_used.')

    add_step(doc, 4, 'Berikan Obat ke Pasien',
             'Beserta: etiket obat (dari label printer terpisah), info signa & aturan pakai, edukasi efek samping.')

    doc.add_heading('5. Cetak Resep', level=1)
    doc.add_paragraph('Detail resep → tombol "🖨️ Cetak Resep". PDF dengan kop RS, list obat, tanda tangan digital dokter, QR verifikasi.')
    add_screenshot(doc, '06-pdf-resep.png', 'Preview PDF resep.')

    doc.add_heading('6. Kelola Stok', level=1)
    doc.add_paragraph('Cek stok aktual: dashboard menampilkan alert obat menipis & mendekati expired.')
    add_screenshot(doc, '06-dashboard-alert-stok.png', 'Dashboard alert stok.')
    add_callout(doc, 'info', 'Fitur input barang masuk & CRUD stok belum ada UI di v1. Hubungi IT untuk tambah stok baru.')

    add_troubleshooting(doc, [
        ('Dispense error "Stok kosong" padahal fisik ada',
         'Kemungkinan: sistem beda dengan aktual (lapor IT), batch expired (cek exp_date), atau di-lock dispense paralel (tunggu 5 detik).'),
        ('Verifikasi resep tidak ada tombolnya', 'TTK_APOTEK tidak bisa verifikasi. Apoteker yang harus.'),
        ('Pasien komplain obat tidak lengkap', 'Cek section batch_used di detail resep. Kalau kosong, dispense belum berhasil.'),
        ('Autocomplete obat lambat', 'Ketik min 3 karakter. Refresh kalau tetap lambat.'),
    ])

    add_contact(doc, [
        ('Konsultasi obat', 'Apoteker Penanggung Jawab (APJ): __________'),
        ('Adverse event / MESO', 'Komite Farmasi & Terapi: __________'),
        ('Bug aplikasi', 'IT Support: __________'),
    ])


# ============================================================
def build_kasir(doc):
    add_cover(doc, 'Panduan Kasir', 'KASIR & KASIR_SUPERVISOR',
              'Panduan untuk role KASIR (kasir.lina) — terima pembayaran, dan KASIR_SUPERVISOR (kasir.ratna) — void pembayaran + export Excel. Prinsip: kasir buat transaksi, koreksi (void) butuh approval supervisor.')
    doc.add_page_break()

    add_toc(doc, ['1. Perbedaan Role', '2. Buka Daftar Tagihan', '3. Terima Pembayaran',
                  '4. Cetak Kuitansi', '5. Cicilan', '6. Void Pembayaran (Supervisor)',
                  '7. Rekap End-of-Shift', '8. Export Excel (Supervisor)',
                  '9. Troubleshooting', '10. Kontak'])
    doc.add_page_break()

    doc.add_heading('1. Perbedaan Role', level=1)
    p = doc.add_paragraph(); p.add_run('KASIR — ').bold = True
    p.add_run('Lihat tagihan, terima pembayaran, cetak kuitansi.')
    p = doc.add_paragraph(); p.add_run('KASIR_SUPERVISOR — ').bold = True
    p.add_run('Semua di atas + void pembayaran (koreksi) + export rekap Excel.')

    doc.add_heading('2. Buka Daftar Tagihan', level=1)
    doc.add_paragraph('Menu "Billing".')
    add_screenshot(doc, '07-billing-list.png', 'Daftar tagihan hari ini.')
    doc.add_paragraph('Filter status: DRAFT / BELUM_LUNAS / CICILAN / LUNAS / KLAIM / VOID')
    add_callout(doc, 'tip', 'Fokus di BELUM_LUNAS dan CICILAN — itu antrian kerja Anda.')

    doc.add_page_break()
    doc.add_heading('3. Terima Pembayaran', level=1)
    add_step(doc, 1, 'Cari Tagihan Pasien', 'Klik detail tagihan.')
    add_screenshot(doc, '07-detail-tagihan.png', 'Detail tagihan lengkap.')

    add_step(doc, 2, 'Klik "💰 Terima Pembayaran"')
    add_screenshot(doc, '07-form-pembayaran.png', 'Form pembayaran.')

    doc.add_paragraph('Isi: metode pembayaran, jumlah (bisa sebagian atau full), referensi eksternal (opsional), catatan.')
    p = doc.add_paragraph(); r = p.add_run('Metode:'); r.bold = True
    add_bullet(doc, 'TUNAI — cash')
    add_bullet(doc, 'DEBIT — kartu debit (EDC)')
    add_bullet(doc, 'KREDIT — kartu kredit')
    add_bullet(doc, 'TRANSFER — bank')
    add_bullet(doc, 'QRIS — GoPay, Dana, ShopeePay')
    add_bullet(doc, 'BPJS / ASURANSI — klaim penjamin')

    add_step(doc, 3, 'Klik Simpan Pembayaran',
             'Sistem hitung sisa. Kalau sisa=0 → LUNAS, kunjungan → SELESAI. Kalau sisa>0 → CICILAN.')
    add_screenshot(doc, '07-pembayaran-berhasil.png', 'Tagihan status LUNAS.')

    doc.add_page_break()
    doc.add_heading('4. Cetak Kuitansi', level=1)
    doc.add_paragraph('Setelah pembayaran, klik "🖨️ Cetak Kuitansi". PDF berisi kop RS, rincian pembayaran, terbilang, QR verifikasi.')
    add_screenshot(doc, '07-pdf-kuitansi.png', 'Preview PDF kuitansi.')
    add_callout(doc, 'tip', 'Print 2 rangkap — 1 untuk pasien, 1 arsip.')

    doc.add_heading('5. Cicilan (Pembayaran Bertahap)', level=1)
    doc.add_paragraph('Terima cicilan pertama seperti biasa. Sistem update ke CICILAN. Saat pasien bayar lagi:')
    add_bullet(doc, 'Cari tagihan yang sama')
    add_bullet(doc, 'Klik Terima Pembayaran lagi')
    add_bullet(doc, 'Sistem otomatis pakai sisa sebagai max jumlah')
    add_screenshot(doc, '07-tagihan-cicilan.png', 'Tagihan status CICILAN dengan history pembayaran.')

    doc.add_page_break()
    doc.add_heading('6. Void Pembayaran (KASIR_SUPERVISOR only)', level=1)
    doc.add_paragraph('Kalau pembayaran salah input, harus di-void oleh supervisor.')
    add_step(doc, 1, 'Detail Tagihan → Cari Pembayaran Salah')
    add_step(doc, 2, 'Klik icon ❌ Void', 'Muncul hanya untuk supervisor.')
    add_step(doc, 3, 'Isi Alasan Void', 'Wajib. Mis. "Salah input jumlah, seharusnya Rp 250.000 bukan Rp 2.500.000".')
    add_callout(doc, 'warning', 'Void tidak menghapus record, hanya tandai is_void=true + reason. Sistem hitung ulang sisa otomatis.')
    add_callout(doc, 'critical', 'Void = audit-heavy. Setiap void tercatat kapan/siapa/alasan. Semua akan diaudit.')

    doc.add_heading('7. Rekap End-of-Shift', level=1)
    doc.add_paragraph('Sebelum ganti shift / tutup kasir:')
    add_bullet(doc, 'Menu Billing → filter tanggal hari ini')
    add_bullet(doc, 'Lihat total tagihan, total dibayar, breakdown metode')
    add_bullet(doc, 'Cocokkan dengan uang tunai (untuk cash), rekening bank (transfer/QRIS/EDC)')
    add_bullet(doc, 'Selisih → cari transaksi mencurigakan (void, cicilan belum tercatat)')

    doc.add_heading('8. Export Excel (KASIR_SUPERVISOR)', level=1)
    doc.add_paragraph('Untuk keuangan bulanan:')
    add_bullet(doc, 'Menu Billing → tombol "📊 Export Excel"')
    add_bullet(doc, 'Filter tanggal + status → Excel download')
    add_screenshot(doc, '07-export-excel.png', 'Tombol export Excel di halaman billing.')

    add_troubleshooting(doc, [
        ('"Simpan Pembayaran" error "Jumlah melebihi sisa"', 'Kolom jumlah max = sisa tagihan. Cek angka.'),
        ('"Simpan Pembayaran" error "Tagihan sudah lunas"', 'Kasir lain sudah bayar. Refresh halaman.'),
        ('Tombol Cetak Kuitansi tidak muncul', 'Hanya muncul setelah pembayaran tersimpan. Refresh.'),
        ('Nomor kuitansi ada "/", apakah normal?', 'Ya. Format PB/2026/09/000123. Sistem replace slash saat generate PDF filename.'),
    ])

    add_contact(doc, [
        ('Konsultasi finansial', 'Bagian Keuangan / Bendahara: __________'),
        ('Approval void', 'Kasir Supervisor / Manajer Keuangan: __________'),
        ('Bug aplikasi', 'IT Support: __________'),
    ])


# ============================================================
def build_igd(doc):
    add_cover(doc, 'Panduan IGD & Triase', 'Petugas IGD (Perawat/Dokter Jaga)',
              'Panduan alur triase IGD. Berbeda dengan poli, prioritas berdasarkan kegawatan (bukan urutan datang). Tidak ada role khusus "IGD" — tugas dilakukan oleh perawat/dokter yang jaga.')
    doc.add_page_break()

    add_toc(doc, ['1. Alur Umum IGD', '2. Kategori Triase',
                  '3. Triase Awal (dalam 5 menit)', '4. Pemeriksaan Dokter Jaga',
                  '5. Admisi ke Rawat Inap dari IGD', '6. Troubleshooting', '7. Kontak'])
    doc.add_page_break()

    doc.add_heading('1. Alur Umum IGD', level=1)
    doc.add_paragraph('1. Registrasi buat kunjungan tipe IGD. 2. Pasien muncul di IGD Board sebagai "BELUM TRIASE" (prioritas tertinggi). 3. Perawat triase dalam 5 menit. 4. Dokter jaga periksa berdasarkan prioritas.')

    doc.add_heading('2. Kategori Triase', level=1)
    doc.add_paragraph('Board menampilkan pasien dalam 5 section (urut prioritas):')
    add_bullet(doc, 'BELUM TRIASE — baru datang, wajib segera dinilai (max 5 menit)', bold_prefix='⚪ ')
    add_bullet(doc, 'MERAH — Resusitasi (henti jantung/napas, syok berat) — SEGERA', bold_prefix='🔴 ')
    add_bullet(doc, 'KUNING — Emergent (< 10 menit)', bold_prefix='🟡 ')
    add_bullet(doc, 'HIJAU — Urgent (< 30 menit)', bold_prefix='🟢 ')
    add_bullet(doc, 'HITAM — DOA (Dead on Arrival) / expectant', bold_prefix='⚫ ')
    add_screenshot(doc, '08-igd-board.png', 'IGD board dengan pasien di 4-5 section.')
    add_callout(doc, 'warning', 'Board auto-refresh 15 detik. Kalau sedang isi form, pause dulu (⏸).')

    doc.add_page_break()
    doc.add_heading('3. Triase Awal (dalam 5 menit)', level=1)
    add_step(doc, 1, 'Cari Pasien di Section "Belum Triase"')
    add_step(doc, 2, 'Klik Tombol "🩺 Triase"')
    add_screenshot(doc, '08-form-triase.png', 'Form triase awal.')

    doc.add_paragraph('Isi:')
    add_bullet(doc, 'Kategori Triase (MERAH/KUNING/HIJAU/HITAM)')
    add_bullet(doc, 'Keluhan Utama')
    add_bullet(doc, 'Tanda Vital: TD, Nadi, Respirasi, Suhu, SpO2, GCS')

    add_step(doc, 3, 'Klik Simpan Triase', 'Pasien pindah dari "Belum Triase" ke section warna kategorinya.')
    add_screenshot(doc, '08-triase-berhasil.png', 'IGD board setelah pasien dikategorikan.')

    doc.add_heading('4. Pemeriksaan Dokter Jaga', level=1)
    doc.add_paragraph('Setelah triase, dokter periksa berdasarkan prioritas (Merah dulu). Alur pemeriksaan sama dengan Rawat Jalan (SOAP + diagnosa) — lihat panduan Dokter. Bedanya:')
    add_bullet(doc, 'Tidak ada nomor antrian (prioritas triase)')
    add_bullet(doc, 'Setelah selesai: pulang (ringan), rawat inap (admisi via DPJP), atau rujuk RS lain')

    doc.add_heading('5. Admisi ke Rawat Inap dari IGD', level=1)
    doc.add_paragraph('Kalau pasien perlu RI: 1. Dokter IGD selesaikan pemeriksaan. 2. Konsul telpon ke DPJP spesialis. 3. DPJP login → menu Rawat Inap → Admisi Baru → pilih pasien kunjungan aktif.')

    add_troubleshooting(doc, [
        ('Pasien tidak sadar, tidak ada identitas', 'Daftar sebagai "Mr. X" / "Mrs. Y" + estimasi umur. Update setelah keluarga datang.'),
        ('MERAH datang bersamaan pasien lain', 'MERAH SELALU DULU. Sisanya tunda. Reinforcement dari poli/on-call.'),
        ('Salah triase (HIJAU seharusnya KUNING)', 'v1 belum ada UI edit triase. Pakai catatan CPPT. IT bisa update DB.'),
        ('Mass casualty event', 'Aktifkan protokol MCI. Sistem handle satu-satu — perlu multiple petugas paralel.'),
    ])

    add_contact(doc, [
        ('Kasus rumit', 'Dokter Konsulen / DPJP on-call: __________'),
        ('Mass casualty', 'Direktur Pelayanan Medis + Kepala IGD: __________'),
        ('Bug aplikasi', 'IT Support: __________'),
    ])


# ============================================================
def build_manager_direksi(doc):
    add_cover(doc, 'Panduan Manager, Direksi & Auditor', 'MANAGER / DIREKSI / AUDITOR / PETUGAS_BPJS',
              'Panduan untuk role manajemen. Manager punya akses master data + Excel. Direksi read-only untuk oversight. Auditor eksklusif akses ke Audit Log untuk compliance.')
    doc.add_page_break()

    add_toc(doc, ['1. Perbandingan Fitur per Role', '2. Alur Manager: Dashboard & Rekap',
                  '3. Alur Direksi: High-Level Overview', '4. Alur Auditor: Audit Log',
                  '5. Alur Petugas BPJS', '6. Troubleshooting', '7. Kontak'])
    doc.add_page_break()

    doc.add_heading('1. Perbandingan Fitur per Role', level=1)
    table = doc.add_table(rows=8, cols=5); table.style = 'Light Grid'
    hdr = ['Fitur', 'MANAGER', 'DIREKSI', 'AUDITOR', 'PETUGAS_BPJS']
    for i, h in enumerate(hdr):
        cell = table.rows[0].cells[i]; cell.text = h
        for p in cell.paragraphs:
            for r in p.runs: r.bold = True
    rows_data = [
        ('Dashboard KPI', 'Y', 'Y', 'Y', 'BPJS only'),
        ('Data Pasien', 'Y', 'Read-only', 'Read-only', 'BPJS only'),
        ('Kunjungan', 'Y', 'Y', 'Y', 'Y'),
        ('Billing / laporan', 'Y', 'Y', 'Y', 'Klaim BPJS'),
        ('Export Excel', 'Y', 'Y', 'Y', 'N'),
        ('Audit Log', 'N', 'N', 'Y', 'N'),
        ('Master Data CRUD', 'Y', 'N', 'N', 'N'),
    ]
    for i, row in enumerate(rows_data, start=1):
        for j, val in enumerate(row):
            table.rows[i].cells[j].text = val

    doc.add_page_break()
    doc.add_heading('2. Alur Manager: Dashboard & Rekap', level=1)
    add_step(doc, 1, 'Dashboard KPI Harian')
    add_screenshot(doc, '09-dashboard-manager.png', 'Dashboard dengan stat cards.')
    doc.add_paragraph('Yang Anda lihat: total pasien, kunjungan RJ/RI/IGD hari ini, occupancy RI, pendapatan hari ini, alert stok obat.')

    add_step(doc, 2, 'Master Data (Update Dokter/Poli/Kamar/Obat)',
             'CRUD via UI belum ada v1. Request ke IT untuk update.')

    add_step(doc, 3, 'Export Rekap untuk Rapat')
    add_screenshot(doc, '09-export-rekap.png', 'Tombol export Excel di halaman billing.')
    doc.add_paragraph('Menu Billing → filter tanggal + status → tombol "📊 Export Excel". Data Pasien juga bisa diexport dari menu Pasien.')

    add_step(doc, 4, 'Monitor Occupancy Real-Time')
    add_screenshot(doc, '09-bed-management-manager.png', 'Bed board manager view.')
    doc.add_paragraph('Target ideal occupancy: 75-85%. Monitor: kamar KOTOR terlalu lama, distribusi per kelas (VIP kosong tapi Kelas III overload?).')

    doc.add_page_break()
    doc.add_heading('3. Alur Direksi: High-Level Overview', level=1)
    doc.add_paragraph('Sama seperti Manager, tapi Direksi biasanya butuh trend pendapatan/BOR/LOS mingguan-bulanan. Trend chart belum ada di v1 — via export Excel + chart manual atau BI tool external.')
    add_callout(doc, 'info', 'Direksi read-only di sistem, tidak input data. Untuk approval kebijakan (mis. tarif naik), coordinate dengan IT + Admin.')

    doc.add_heading('4. Alur Auditor: Audit Log', level=1)
    doc.add_paragraph('Menu "Audit Log" (di sidebar section "KEPATUHAN").')
    add_screenshot(doc, '09-audit-log-viewer.png', 'Halaman audit log dengan filter.')
    doc.add_paragraph('Yang Anda bisa lihat:')
    add_bullet(doc, 'Siapa (causer) mengubah data')
    add_bullet(doc, 'Kapan (timestamp)')
    add_bullet(doc, 'Data apa (Pasien / Kunjungan / User)')
    add_bullet(doc, 'Aksi apa (created / updated / deleted)')
    add_bullet(doc, 'Perubahan detail (klik "Lihat detail" untuk JSON)')

    p = doc.add_paragraph(); r = p.add_run('Skenario Audit Umum:'); r.bold = True
    add_bullet(doc, 'Pasien komplain data diubah tanpa sepengetahuan: filter Subject Type = Pasien, subject_id = UUID pasien. Lihat siapa & kapan.')
    add_bullet(doc, 'Kasir dicurigai manipulasi: filter Subject = Pembayaran, Causer = kasir tsb, Aksi = updated/deleted. Cek void reasons.')
    add_bullet(doc, 'Dispute tagihan: filter Subject = Tagihan. Cek history perubahan status.')
    add_callout(doc, 'tip', 'Untuk audit routine, filter "deleted" — data yang dihapus paling sering jadi bahan pertanyaan.')

    doc.add_page_break()
    doc.add_heading('5. Alur Petugas BPJS', level=1)
    add_callout(doc, 'info',
        'Fitur klaim BPJS belum ter-implement full di v1 — bridging V-Claim direncanakan Tier 3. Sementara alur manual: lihat kunjungan penjamin BPJS via menu Billing → cetak dokumen manual → input ke aplikasi V-Claim BPJS terpisah.')

    add_troubleshooting(doc, [
        ('Dashboard tidak update pendapatan hari ini', 'Cache 5 menit. Refresh halaman.'),
        ('Export Excel kosong', 'Filter terlalu ketat. Longgarkan filter.'),
        ('Audit log tidak muncul aktivitas terbaru', 'Cek filter tanggal. Longgarkan ke "7 hari lalu".'),
        ('Saya Direksi tapi tidak bisa lihat detail resep pasien', 'Sengaja. Direksi tidak akses PHI detail — hanya statistik agregat.'),
    ])

    add_contact(doc, [
        ('Kebijakan strategis', 'Direktur Utama: __________'),
        ('Data operasional', 'Manager Pelayanan / Manager Keuangan: __________'),
        ('Audit finding', 'Komite Audit Internal: __________'),
        ('Bug aplikasi', 'IT Support: __________'),
    ])


# ============================================================
def build_admin(doc):
    add_cover(doc, 'Panduan Administrator IT', 'SUPER_ADMIN',
              'Panduan untuk role SUPER_ADMIN (contoh username: admin). Akses ke semua fitur, bypass semua permission check. Karena akses besar, audit log memantau ketat — jangan sharing akun.')
    doc.add_page_break()

    add_toc(doc, ['1. Peran & Tanggung Jawab', '2. Setup Awal Pasca Go-Live',
                  '3. Manage User', '4. Monitor Sistem',
                  '5. Manage Master Data (via Tinker)', '6. Backup & Restore',
                  '7. Response Insiden', '8. Update Aplikasi',
                  '9. Troubleshooting', '10. Kontak'])
    doc.add_page_break()

    doc.add_heading('1. Peran & Tanggung Jawab', level=1)
    add_bullet(doc, 'Manage user (buat, aktifkan/nonaktifkan, reset password)')
    add_bullet(doc, 'Manage master data (dokter, poli, kamar, obat)')
    add_bullet(doc, 'Monitor sistem (health check, log, backup)')
    add_bullet(doc, 'Response ke insiden (bug, data corruption, security)')
    add_bullet(doc, 'Koordinasi dengan vendor developer')
    add_callout(doc, 'critical', 'Audit log memantau aktivitas Anda ketat. Setiap perubahan tercatat atas nama username Anda. Jangan sharing akun.')

    doc.add_heading('2. Setup Awal Pasca Go-Live', level=1)
    doc.add_heading('Ganti Password Default', level=2)
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run('php artisan tinker\n>>> User::where(\'username\', \'admin\')->first()->update([\'password\' => Hash::make(\'PASSWORD_KUAT_BARU\')]);')
    r.font.name = 'Consolas'; r.font.size = Pt(9)
    add_callout(doc, 'warning', 'Password kuat min 16 karakter, mix huruf besar/kecil/angka/simbol.')

    doc.add_heading('Nonaktifkan User Default Yang Tidak Dipakai', level=2)
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run('>>> User::where(\'username\', \'dr.iqbal\')->update([\'is_active\' => false]);')
    r.font.name = 'Consolas'; r.font.size = Pt(9)

    doc.add_heading('Buat User Real Sesuai Staff Aktual', level=2)
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5)
    code = '''>>> $u = User::create([
...     'username' => 'dr.budi.internist',
...     'name' => 'dr. Budi Susanto, Sp.PD',
...     'email' => 'budi@rs.namars.id',
...     'password' => Hash::make('temp_password'),
...     'is_active' => true,
...     'email_verified_at' => now(),
... ]);
>>> $u->assignRole('DOKTER_SPESIALIS');'''
    r = p.add_run(code); r.font.name = 'Consolas'; r.font.size = Pt(9)

    doc.add_page_break()
    doc.add_heading('3. Manage User', level=1)
    doc.add_heading('Reset Password User', level=2)
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run('>>> User::where(\'username\', \'apt.fitri\')->update([\'password\' => Hash::make(\'password_baru\')]);')
    r.font.name = 'Consolas'; r.font.size = Pt(9)
    doc.add_paragraph('Sampaikan password baru via kanal aman (Signal / kertas tertutup) — bukan WA/email biasa.')

    doc.add_heading('Nonaktifkan (resign)', level=2)
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run('>>> User::where(\'username\', \'kasir.mawar\')->update([\'is_active\' => false]);')
    r.font.name = 'Consolas'; r.font.size = Pt(9)
    doc.add_paragraph('User yang is_active=false otomatis logout kalau sedang session. Data yang pernah mereka input tetap ada.')

    doc.add_heading('4. Monitor Sistem', level=1)
    doc.add_heading('Health Check Endpoint', level=2)
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run('curl https://sihrs.namars.id/up/health'); r.font.name = 'Consolas'; r.font.size = Pt(9)
    doc.add_paragraph('Output JSON dengan checks: database, cache, storage, disk_space, backup. Kalau "status": "unhealthy" → HTTP 503. Segera investigate.')

    doc.add_heading('Sentry (Error Monitoring)', level=2)
    doc.add_paragraph('Login ke sentry.io atau URL self-host. Filter environment = production. Cek: unresolved errors, frequency, affected users. Prioritaskan bug: frekuensi >100, multi-user, patient-safety.')

    doc.add_heading('Log Files', level=2)
    add_bullet(doc, 'storage/logs/laravel.log — Laravel default (rotate daily)')
    add_bullet(doc, 'storage/logs/audit.log — audit trail (retensi 5 tahun)')
    add_bullet(doc, 'storage/logs/bpjs.log — komunikasi BPJS')
    add_bullet(doc, 'storage/logs/satusehat.log — komunikasi SATUSEHAT')
    add_bullet(doc, 'storage/logs/worker.log — queue worker')
    add_bullet(doc, 'storage/logs/backup.log — job backup harian')

    doc.add_page_break()
    doc.add_heading('5. Manage Master Data (via Tinker)', level=1)
    doc.add_paragraph('UI CRUD belum ada v1 — via CLI. Contoh tambah dokter baru:')
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5)
    code = '''>>> $dokter = Dokter::create([
...     'kode' => 'DR00021',
...     'sip' => '445/000/DKK/2026',
...     'nama' => 'Ahmad Fadhil',
...     'gelar_depan' => 'dr.', 'gelar_belakang' => 'Sp.B',
...     'spesialisasi' => 'Bedah',
...     'jasa_konsul' => 200000,
...     'is_active' => true,
... ]);'''
    r = p.add_run(code); r.font.name = 'Consolas'; r.font.size = Pt(9)

    doc.add_heading('6. Backup & Restore', level=1)
    doc.add_heading('Backup Manual', level=2)
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run('sudo -u www-data php artisan sihrs:backup'); r.font.name = 'Consolas'; r.font.size = Pt(9)

    doc.add_heading('Restore Backup', level=2)
    add_callout(doc, 'critical', 'DANGER: Restore = HAPUS DATA SAAT INI. Jangan restore di production tanpa persetujuan direksi.')
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5)
    code = '''# 1. Backup dulu DB current (safety net)
sudo -u www-data php artisan sihrs:backup

# 2. Restore
gunzip -c /var/backups/sihrs/sihrs_YYYYMMDD.sql.gz | mysql -u sihrs_app -p sihrs

# 3. Verify
mysql -u sihrs_app -p sihrs -e "SELECT COUNT(*) FROM pasien"'''
    r = p.add_run(code); r.font.name = 'Consolas'; r.font.size = Pt(9)

    doc.add_page_break()
    doc.add_heading('7. Response Insiden', level=1)
    doc.add_heading('Data Pasien Corrupt', level=2)
    doc.add_paragraph('1. Screenshot masalah UI. 2. Buka Audit Log, cek aktivitas delete/update pada waktu tsb. 3. Malicious → identifikasi user, nonaktifkan. 4. Bug software → restore dari backup terdekat (setelah persetujuan direksi). 5. Report ke vendor.')

    doc.add_heading('Aplikasi Down', level=2)
    doc.add_paragraph('1. Cek /up/health return apa. 2. Cek server: systemctl status nginx php8.2-fpm mysql. 3. Cek log: tail -100 storage/logs/laravel.log. 4. Restart service yang mati. 5. Escalate ke vendor.')

    doc.add_heading('Data Breach (Dicurigai)', level=2)
    doc.add_paragraph('1. Isolate — matikan akses external. 2. Preserve evidence — jangan restart, jangan hapus log. 3. Log semua aktivitas 24 jam terakhir. 4. Report ke Direksi + tim legal + Kominfo (kewajiban UU 27/2022 dalam 72 jam). 5. Post-mortem.')

    doc.add_heading('8. Update Aplikasi (Deploy Baru)', level=1)
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5)
    code = '''cd /var/www/sihrs
sudo -u www-data php artisan down --render="errors::503"
sudo -u www-data git pull origin main
sudo -u www-data composer install --no-dev --optimize-autoloader
sudo -u www-data npm ci && sudo -u www-data npm run build
sudo -u www-data php artisan migrate --force
sudo -u www-data php artisan config:cache
sudo -u www-data php artisan route:cache
sudo -u www-data php artisan view:cache
sudo systemctl restart php8.2-fpm
sudo supervisorctl restart sihrs-worker:*
sudo -u www-data php artisan up'''
    r = p.add_run(code); r.font.name = 'Consolas'; r.font.size = Pt(9)
    doc.add_paragraph('Downtime target: < 2 menit kalau tidak ada migrasi berat.')

    add_troubleshooting(doc, [
        ('Staff request akses fitur yang tidak ada di role', 'Jangan langsung berikan. Diskusi dengan Manager unit — mungkin butuh review permission struktur.'),
        ('Sistem lambat jam sibuk', 'Cek /up/health — kalau DB latency > 100ms, enable slow_query_log. Cari query lambat.'),
        ('User mengeluh session logout tiba-tiba', 'Session lifetime = 2 jam idle. Atau akun dinonaktifkan → middleware auto-logout.'),
        ('BPJS bridging kapan aktif', 'Tier 3 roadmap. Sementara klaim manual.'),
    ])

    add_contact(doc, [
        ('Vendor developer', '[Nama Vendor] — __________'),
        ('Escalation', 'CTO / Kepala IT: __________'),
        ('Emergency 24/7', 'On-call: __________'),
    ])


# ============================================================
# MAIN
# ============================================================

BUILDERS = {
    '00-getting-started': ('00-Getting-Started', build_getting_started),
    '01-registrasi': ('01-Panduan-Registrasi', build_registrasi),
    '02-dokter': ('02-Panduan-Dokter', build_dokter),
    '03-perawat': ('03-Panduan-Perawat', build_perawat),
    '04-lab': ('04-Panduan-Laboratorium', build_lab),
    '05-radiologi': ('05-Panduan-Radiologi', build_radiologi),
    '06-farmasi': ('06-Panduan-Farmasi', build_farmasi),
    '07-kasir': ('07-Panduan-Kasir', build_kasir),
    '08-igd': ('08-Panduan-IGD', build_igd),
    '09-manager-direksi': ('09-Panduan-Manager-Direksi', build_manager_direksi),
    '10-admin': ('10-Panduan-Administrator', build_admin),
}


def build(key):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2); section.bottom_margin = Cm(2)
        section.left_margin = Cm(2); section.right_margin = Cm(2)
    set_document_style(doc)

    if key not in BUILDERS:
        print(f'[ERR] Role tidak dikenal: {key}. Available: {list(BUILDERS.keys())}')
        sys.exit(1)

    filename, builder = BUILDERS[key]
    builder(doc)

    out_file = OUT_DIR / f'{filename}.docx'
    doc.save(out_file)
    print(f'[OK] {out_file.name} ({out_file.stat().st_size // 1024} KB)')


if __name__ == '__main__':
    arg = sys.argv[1] if len(sys.argv) > 1 else 'all'
    if arg == 'all':
        for k in BUILDERS.keys():
            build(k)
    else:
        # Support pass 'registrasi' → maps to '01-registrasi'
        for k in BUILDERS.keys():
            if arg in k:
                build(k); sys.exit(0)
        print(f'[ERR] Not found: {arg}')
